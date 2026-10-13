"""Insurance and TPA claim tracking, kept alongside the patient bills.

Gated exactly like the patient bills (has_kickbacks_access), so the same
plan and manager-page rules apply.
"""
import datetime

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.generic import View

from . import services
from .access import has_kickbacks_access
from .forms import ClaimPaymentForm, InsuranceClaimForm, PatientPaymentForm
from .models import ClaimPayment, InsuranceClaim, KickbackEntry, PaymentMode, SalesEntry, sync_claim_totals
from .services import _patient_label
from .views import TenantLoginRequiredMixin, _render_statement_pdf

# An open claim older than this many days since it was submitted is flagged
# for follow-up with the insurer or TPA.
FOLLOW_UP_DAYS = 30

PAGE_SIZE = 50

# Filter values that are not a single status.
FILTER_OPEN = "open"
FILTER_OVERDUE = "overdue"


def _book_revenue(claim, payment, email: str) -> None:
    """Money received on a claim is revenue on the day it arrives: a Daily Book
    entry linked to the claim's bill, so the bill's received total follows.
    Deductions are not revenue and are not booked."""
    if payment.amount <= 0:
        return
    entry = SalesEntry.objects.create(
        date=payment.date,
        amount=payment.amount,
        payment_mode=PaymentMode.BANK,
        channel=services.get_finance_settings().ip_revenue_category,
        kickback_bill=claim.bill,
        note=f"Insurance claim {claim.claim_number or 'no ref'} · {claim.insurer_display}"[:255],
        created_by_email=email,
    )
    payment.sales_entry = entry
    payment.save()


def _remove_payment(payment) -> int:
    """Deletes a payment and the revenue entry booked for it. Returns its claim id."""
    claim_id = payment.claim_id
    if payment.sales_entry_id:
        entry = SalesEntry.objects.filter(pk=payment.sales_entry_id).first()
        if entry is not None:
            entry.delete()
    payment.delete()
    return claim_id


def _is_overdue(claim, today) -> bool:
    days = claim.days_since_submitted(today)
    return claim.is_open and days is not None and days > FOLLOW_UP_DAYS


def _summary(claims, today) -> dict:
    zero = services.ZERO
    open_claims = [c for c in claims if c.is_open]
    return {
        "claimed": sum((c.claimed_amount for c in claims), zero),
        "approved": sum((c.approved_amount for c in claims), zero),
        "discount": sum((c.discount_amount for c in claims), zero),
        "patient": sum((c.patient_share for c in claims), zero),
        "received": sum((c.received_amount for c in claims), zero),
        "deducted": sum((c.deduction_amount for c in claims), zero),
        "outstanding": sum((c.outstanding_amount for c in open_claims), zero),
        "open_count": len(open_claims),
        "overdue_count": sum(1 for c in open_claims if _is_overdue(c, today)),
    }


def _apply_filters(claims, query, status, today):
    """The list's search and status filter, shared with the exports."""
    if query:
        needle = query.lower()
        claims = [c for c in claims if needle in " ".join([
            c.claim_number, c.patient_name, c.insurer_name, c.tpa_name, c.policy_number,
            c.bill.bill_number if c.bill else "",
        ]).lower()]
    if status == FILTER_OPEN:
        claims = [c for c in claims if c.is_open]
    elif status == FILTER_OVERDUE:
        claims = [c for c in claims if _is_overdue(c, today)]
    elif status:
        claims = [c for c in claims if c.status == status]
    return claims


def _export_claims(request):
    """The claims the list shows for the current filters, for the exports."""
    today = datetime.date.today()
    query = (request.GET.get("q") or "").strip()
    status = (request.GET.get("status") or "").strip()
    return _apply_filters(list(InsuranceClaim.objects.all()), query, status, today), today


class InsuranceClaimsView(TenantLoginRequiredMixin, View):
    template_name = "finance/claims.html"

    def get(self, request, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        today = datetime.date.today()
        query = (request.GET.get("q") or "").strip()
        status = (request.GET.get("status") or "").strip()

        all_claims = list(InsuranceClaim.objects.all())
        summary = _summary(all_claims, today)
        claims = _apply_filters(all_claims, query, status, today)

        page = Paginator(claims, PAGE_SIZE).get_page(request.GET.get("page"))
        for claim in page.object_list:
            claim.days_open = claim.days_since_submitted(today)
            claim.overdue = _is_overdue(claim, today)

        filters = request.GET.copy()
        filters.pop("page", None)
        return render(request, self.template_name, {
            "active_nav": "claims",
            "organization": request.user.organization,
            "page": page,
            "claims_count": len(claims),
            "summary": summary,
            "filter_query": query,
            "filter_status": status,
            "status_choices": InsuranceClaim.Status.choices,
            "open_filter": FILTER_OPEN,
            "overdue_filter": FILTER_OVERDUE,
            "follow_up_days": FOLLOW_UP_DAYS,
            "page_query": ("?" + filters.urlencode()) if filters else "",
            "export_query": filters.urlencode(),
        })


def _bill_amounts(bills) -> dict:
    """The amounts shown under the Patient bill dropdown, keyed by bill id, so
    the page can update them when another bill is picked."""
    def money(value):
        return f"{value:,.2f}"
    return {
        str(bill.pk): {
            "final": money(bill.final_amount),
            "final_raw": str(bill.final_amount),
            "insurer": money(bill.insurer_share),
            "patient": money(bill.patient_amount),
        }
        for bill in bills
    }


def _render_form(request, form, claim, *, title):
    return render(request, "finance/claim_form.html", {
        "active_nav": "claims",
        "organization": request.user.organization,
        "form": form,
        "claim": claim,
        "title": title,
        "bill_amounts": _bill_amounts(form.fields["bill"].queryset),
    })


class AddInsuranceClaimView(TenantLoginRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        initial = {"claim_type": InsuranceClaim.ClaimType.CLAIM, "status": InsuranceClaim.Status.DRAFT}
        # "Add claim" from a patient bill pre-fills what the bill already knows.
        bill = None
        if request.GET.get("bill"):
            try:
                bill = KickbackEntry.objects.filter(pk=request.GET["bill"]).first()
            except (ValueError, ValidationError):
                bill = None  # a mangled id in the URL just means a blank form
        if bill is not None:
            initial.update({
                "bill": bill.pk,
                "insurer_name": bill.insurer_name,
                "policy_number": bill.policy_number,
                "claimed_amount": bill.insurer_share,
            })
            if bill.preauth_number:
                initial["claim_number"] = bill.preauth_number
        form = InsuranceClaimForm(initial=initial)
        return _render_form(request, form, None, title="New insurance claim")

    def post(self, request, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        form = InsuranceClaimForm(request.POST)
        if not form.is_valid():
            messages.error(request, "Couldn't save that claim — check the highlighted fields.")
            return _render_form(request, form, None, title="New insurance claim")
        claim = form.save(commit=False)
        claim.created_by_email = request.user.email
        claim.save()
        form.apply_patient_share(claim)
        messages.success(request, "Claim added.")
        return redirect("finance:claim_detail", pk=claim.pk)


class EditInsuranceClaimView(TenantLoginRequiredMixin, View):
    def get(self, request, pk, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        claim = get_object_or_404(InsuranceClaim, pk=pk)
        return _render_form(request, InsuranceClaimForm(instance=claim), claim, title="Edit insurance claim")

    def post(self, request, pk, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        claim = get_object_or_404(InsuranceClaim, pk=pk)
        form = InsuranceClaimForm(request.POST, instance=claim)
        if not form.is_valid():
            messages.error(request, "Couldn't save that claim — check the highlighted fields.")
            return _render_form(request, form, claim, title="Edit insurance claim")
        form.save()
        sync_claim_totals(claim.pk)  # a status edit can settle or reopen the claim
        messages.success(request, "Claim updated.")
        return redirect("finance:claim_detail", pk=claim.pk)


def _render_detail(request, claim, payment_form):
    today = datetime.date.today()
    payments = sorted(ClaimPayment.objects.filter(claim=claim), key=lambda p: (p.date, p.created_at))
    return render(request, "finance/claim_detail.html", {
        "active_nav": "claims",
        "organization": request.user.organization,
        "claim": claim,
        "payments": payments,
        "payment_form": payment_form,
        "days_open": claim.days_since_submitted(today),
        "overdue": _is_overdue(claim, today),
        "follow_up_days": FOLLOW_UP_DAYS,
    })


class InsuranceClaimDetailView(TenantLoginRequiredMixin, View):
    def get(self, request, pk, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        claim = get_object_or_404(InsuranceClaim, pk=pk)
        return _render_detail(request, claim, ClaimPaymentForm(claim=claim))


class AddClaimPaymentView(TenantLoginRequiredMixin, View):
    def post(self, request, pk, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        claim = get_object_or_404(InsuranceClaim, pk=pk)
        form = ClaimPaymentForm(request.POST, claim=claim)
        if not form.is_valid():
            messages.error(request, "Payment not saved — see the message under the form.")
            return _render_detail(request, claim, form)
        payment = form.save(commit=False)
        payment.claim = claim
        payment.save()
        _book_revenue(claim, payment, request.user.email)
        sync_claim_totals(claim.pk)
        messages.success(request, "Payment recorded and added to revenue.")
        return redirect("finance:claim_detail", pk=claim.pk)


class AddPatientPaymentView(TenantLoginRequiredMixin, View):
    """Money the patient pays towards their share. It goes to the Daily Book
    as revenue on the claim's admission, and is not counted against the claim."""

    def post(self, request, pk, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        claim = get_object_or_404(InsuranceClaim, pk=pk)
        if claim.bill is None:
            messages.error(request, "Link this claim to an IP admission before recording a patient payment.")
            return redirect("finance:claim_detail", pk=claim.pk)
        form = PatientPaymentForm(request.POST)
        if not form.is_valid():
            messages.error(request, "Patient payment not saved — check the amount and date.")
            return redirect("finance:claim_detail", pk=claim.pk)
        data = form.cleaned_data
        bill = claim.bill
        note = f"Patient payment · {_patient_label(bill)}"
        if data["note"]:
            note = f"{note} · {data['note']}"
        SalesEntry.objects.create(
            date=data["date"],
            amount=data["amount"],
            payment_mode=data["payment_mode"],
            channel=services.get_finance_settings().ip_revenue_category,
            kickback_bill=bill,
            note=note[:255],
            created_by_email=request.user.email,
        )
        messages.success(request, "Patient payment recorded and added to revenue.")
        return redirect("finance:claim_detail", pk=claim.pk)


class DeleteClaimPaymentView(TenantLoginRequiredMixin, View):
    def post(self, request, pk, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        payment = get_object_or_404(ClaimPayment, pk=pk)
        claim_id = _remove_payment(payment)
        sync_claim_totals(claim_id)
        messages.success(request, "Payment removed.")
        return redirect("finance:claim_detail", pk=claim_id)


class DeleteInsuranceClaimView(TenantLoginRequiredMixin, View):
    def post(self, request, pk, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        claim = get_object_or_404(InsuranceClaim, pk=pk)
        # Payments are removed explicitly: a Google Sheets-backed org does not
        # cascade deletes the way the database does.
        for payment in ClaimPayment.objects.filter(claim=claim):
            _remove_payment(payment)
        claim.delete()
        messages.success(request, "Claim deleted.")
        return redirect("finance:claims")


CLAIM_EXPORT_HEADERS = [
    "Claim / reference no.", "Type", "Patient", "Invoice no.", "Insurer", "TPA", "Policy / member ID",
    "Status", "Claimed", "Approved", "Received", "Deducted", "Outstanding",
    "Submitted on", "Decision date", "Last payment", "Days since submitted", "Remarks",
]
CLAIM_EXPORT_MONEY = ("Claimed", "Approved", "Received", "Deducted", "Outstanding")


def _claim_export_row(claim, today):
    return [
        claim.claim_number, claim.get_claim_type_display(), claim.patient_name,
        claim.bill.bill_number if claim.bill else "", claim.insurer_name, claim.tpa_name,
        claim.policy_number, claim.get_status_display(),
        claim.claimed_amount, claim.approved_amount, claim.received_amount,
        claim.deduction_amount, claim.outstanding_amount,
        claim.submitted_date.isoformat() if claim.submitted_date else "",
        claim.decision_date.isoformat() if claim.decision_date else "",
        claim.last_payment_date.isoformat() if claim.last_payment_date else "",
        claim.days_since_submitted(today) if claim.submitted_date else "",
        claim.remarks,
    ]


class InsuranceClaimsExcelView(TenantLoginRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        claims, today = _export_claims(request)
        wb = Workbook()
        ws = wb.active
        ws.title = "Insurance Claims"
        ws.append(CLAIM_EXPORT_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        money_columns = [i + 1 for i, h in enumerate(CLAIM_EXPORT_HEADERS) if h in CLAIM_EXPORT_MONEY]
        for claim in claims:
            ws.append(_claim_export_row(claim, today))
        for column in money_columns:
            for row in ws.iter_rows(min_row=2, min_col=column, max_col=column):
                for cell in row:
                    cell.number_format = "#,##0.00"
        for index, header in enumerate(CLAIM_EXPORT_HEADERS, start=1):
            ws.column_dimensions[get_column_letter(index)].width = max(12, min(40, len(header) + 4))
        ws.freeze_panes = "A2"

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="insurance-claims.xlsx"'
        wb.save(response)
        return response


class InsuranceClaimsPdfView(TenantLoginRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        if not has_kickbacks_access(request):
            return redirect("finance:dashboard")
        claims, today = _export_claims(request)
        for claim in claims:
            claim.days_open = claim.days_since_submitted(today)
        totals = {
            "claimed": sum((c.claimed_amount for c in claims), services.ZERO),
            "approved": sum((c.approved_amount for c in claims), services.ZERO),
            "received": sum((c.received_amount for c in claims), services.ZERO),
            "deducted": sum((c.deduction_amount for c in claims), services.ZERO),
            "outstanding": sum((c.outstanding_amount for c in claims if c.is_open), services.ZERO),
        }
        context = {
            "organization": request.user.organization,
            "claims": claims,
            "totals": totals,
            "filter_query": (request.GET.get("q") or "").strip(),
            "filter_status": (request.GET.get("status") or "").strip(),
            "generated_on": today,
        }
        return _render_statement_pdf(
            request, "finance/pdf/claims_pdf.html", context, "insurance-claims.pdf",
            fallback_url_name="finance:claims",
        )
