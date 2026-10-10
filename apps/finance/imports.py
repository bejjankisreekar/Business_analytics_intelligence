"""CSV/Excel import for Revenue entries — read entirely in memory, never
written to disk or the database until the user reviews and explicitly saves
(see views.ImportSalesUploadView / ImportSalesCommitView). Kept separate
from services.py since this is file parsing/IO, not analytics.
"""
import csv
import datetime
import io
import re
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from .models import BankAccount, Category, Customer, PaymentMode, Subcategory

MAX_IMPORT_ROWS = 500
MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB

# A small, realistic example file for the upload page's "download sample" link —
# every column name here matches a TARGET_FIELDS label exactly, so a client
# can literally fill this in and upload it back unchanged.
SAMPLE_ROWS = [
    ["Date", "Category", "Sub-category", "Customer", "Qty / Count", "Gross Amount", "Discount", "Payment Mode", "Bank", "Note"],
    ["2026-09-28", "Sales", "Walk-in", "Walk-in Customer", "1", "500", "0", "Cash", "", ""],
    ["2026-09-28", "Online Orders", "", "Walk-in Customer", "3", "1200", "100", "Cash", "", "Loyalty discount"],
    ["2026-09-29", "Service Revenue", "", "Walk-in Customer", "2", "2000", "0", "Bank", "", ""],
]

# (field name, label shown in the mapping UI, required)
TARGET_FIELDS = [
    ("date", "Date", True),
    ("channel", "Category", False),
    ("subcategory", "Sub-category", False),
    ("customer", "Customer", False),
    ("quantity", "Qty / Count", False),
    ("gross_amount", "Gross Amount", True),
    ("discount", "Discount", False),
    ("payment_mode", "Payment Mode", False),
    ("bank_account", "Bank", False),
    ("note", "Note", False),
]

_FIELD_ALIASES = {
    "date": ["date", "txn date", "transaction date", "bill date", "entry date", "invoice date"],
    "channel": ["category", "channel", "revenue channel", "department", "service"],
    "subcategory": ["sub category", "subcategory", "sub-category", "detail", "item", "service detail"],
    "customer": ["customer", "patient", "patient name", "client", "payer name"],
    "quantity": ["qty", "quantity", "count", "units", "no of patients", "number of patients"],
    "gross_amount": ["gross amount", "gross", "billed amount", "bill amount", "total amount", "amount"],
    "discount": ["discount", "discount amount", "concession", "concession amount"],
    "payment_mode": ["payment mode", "mode of payment", "mode", "payment type", "via"],
    "bank_account": ["bank", "bank account", "account"],
    "note": ["note", "notes", "remark", "remarks", "description"],
}


class ImportParseError(ValueError):
    """A problem with the uploaded file itself — shown to the user as-is."""


def _normalize(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def parse_upload(uploaded_file) -> tuple[list[str], list[list]]:
    """Read an uploaded .csv/.xlsx file fully into memory and return
    (headers, data_rows). The file is never saved — only its bytes are read,
    right here, for this one request."""
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        raise ImportParseError(f"That file is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB — split it into smaller files.")

    name = (uploaded_file.name or "").lower()
    if name.endswith(".csv"):
        text = uploaded_file.read().decode("utf-8-sig", errors="replace")
        raw_rows = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        workbook = load_workbook(filename=io.BytesIO(uploaded_file.read()), read_only=True, data_only=True)
        sheet = workbook.active
        raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        raise ImportParseError("Unsupported file type — please upload a .csv or .xlsx file.")

    raw_rows = [row for row in raw_rows if any(str(c).strip() for c in row if c is not None)]
    if not raw_rows:
        raise ImportParseError("That file has no data in it.")

    headers = [str(c).strip() if c is not None else "" for c in raw_rows[0]]
    data_rows = raw_rows[1:]
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise ImportParseError(
            f"That file has {len(data_rows)} rows — imports are capped at {MAX_IMPORT_ROWS} rows "
            "at a time. Split it into smaller files."
        )

    width = len(headers)
    normalized_rows = []
    for row in data_rows:
        row = list(row[:width]) + [""] * max(0, width - len(row))
        normalized_rows.append(["" if c is None else c for c in row])
    return headers, normalized_rows


def guess_column_mapping(headers: list[str]) -> dict[str, str | None]:
    """Best-effort guess of which spreadsheet header feeds which target
    field, by normalized exact match first, then loose containment. Every
    header is used for at most one field."""
    normalized = {h: _normalize(h) for h in headers}
    used: set[str] = set()
    mapping: dict[str, str | None] = {}
    for field, _label, _required in TARGET_FIELDS:
        aliases = [_normalize(a) for a in _FIELD_ALIASES[field]]
        match = next((h for h in headers if h not in used and normalized[h] in aliases), None)
        if not match:
            match = next(
                (
                    h for h in headers if h not in used and normalized[h]
                    and any(a and (a in normalized[h] or normalized[h] in a) for a in aliases)
                ),
                None,
            )
        mapping[field] = match
        if match:
            used.add(match)
    return mapping


def _parse_date_cell(value) -> datetime.date | None:
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value or "").strip().split("T")[0]
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%y", "%d-%b-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal_cell(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    if not cleaned:
        return Decimal("0")
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return Decimal("0")


def _parse_int_cell(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def _match_by_name(value, objects) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    for obj in objects:
        if obj.name.strip().lower() == text:
            return str(obj.pk)
    return ""


def _match_payment_mode(value) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    return PaymentMode.CASH if "cash" in text else PaymentMode.BANK


KICKBACK_SAMPLE_ROWS = [
    ["Bill Number", "Patient Name", "Address", "Admission Date", "Discharge Date", "Consult Doctor", "Reffered By",
     "Final Amount"],
    ["BL-1001", "Ramesh Kumar", "Kothapalli, Narasapur, West Godavari", "2026-10-01", "2026-10-05",
     "Dr. Akash Bansal", "Dr. Aisha Khan", "60000"],
    ["BL-1002", "Sita Devi", "Madhapur, Hyderabad", "2026-10-03", "", "Dr. Neha Verma", "", "35000"],
]

# Every uploaded column is matched against these aliases by normalized name
# (exact, then loose containment) — a column that matches nothing here is
# simply never read, i.e. ignored, so an uploaded file can carry extra
# columns (hospital internal refs, etc.) without issue. "Net Amount" is
# deliberately not mappable — KickbackEntry.save() always derives it.
KICKBACK_FIELD_ALIASES = {
    "bill_number": ["bill number", "bill no", "billno", "invoice number", "invoice no"],
    "patient_name": ["patient name", "patient", "name of patient", "pt name", "patient full name"],
    "patient_address": ["address", "patient address", "addr", "residence"],
    "admission_date": ["admission date", "admit date", "date of admission"],
    "discharge_date": ["discharge date", "date of discharge"],
    "consulting_doctor": ["consult doctor", "consulting doctor", "doctor", "consultant"],
    "referred_by": ["reffered by", "referred by", "reference", "referral", "referred"],
    "final_amount": ["final amount", "bill amount", "total amount", "total bill amount"],
}

KICKBACK_REQUIRED_FIELDS = ("bill_number", "admission_date")


def map_kickback_columns(headers: list[str]) -> dict[str, str | None]:
    """Best-effort match of uploaded headers to KickbackEntry fields — see
    KICKBACK_FIELD_ALIASES. Any header that doesn't match a known column
    name is left out of the mapping and never read."""
    normalized = {h: _normalize(h) for h in headers}
    used: set[str] = set()
    mapping: dict[str, str | None] = {}
    for field, aliases in KICKBACK_FIELD_ALIASES.items():
        alias_norms = [_normalize(a) for a in aliases]
        mapping[field] = next((h for h in headers if h not in used and normalized[h] in alias_norms), None)
        if mapping[field]:
            used.add(mapping[field])
    for field, aliases in KICKBACK_FIELD_ALIASES.items():
        if mapping[field]:
            continue
        alias_norms = [_normalize(a) for a in aliases]
        mapping[field] = next(
            (
                h for h in headers if h not in used and normalized[h]
                and any(a and (a in normalized[h] or normalized[h] in a) for a in alias_norms)
            ),
            None,
        )
        if mapping[field]:
            used.add(mapping[field])
    return mapping


def build_kickback_rows(
    headers: list[str], data_rows: list[list], mapping: dict[str, str | None]
) -> tuple[list[dict], list[tuple[int, str]]]:
    """Applies `mapping` to every parsed row, producing one
    KickbackEntry-ready dict per valid row, plus (row_number, reason) for
    every row skipped outright (missing bill number or an unrecognized
    admission date) — row_number counts the uploaded file's own rows,
    header included, so it matches what the user sees in a spreadsheet."""
    index = {h: i for i, h in enumerate(headers)}

    def cell(row, field):
        header = mapping.get(field)
        if not header or header not in index:
            return ""
        i = index[header]
        return row[i] if i < len(row) else ""

    rows = []
    errors = []
    for row_number, row in enumerate(data_rows, start=2):
        bill_number = str(cell(row, "bill_number") or "").strip()
        admission_date = _parse_date_cell(cell(row, "admission_date"))
        row_errors = []
        if not bill_number:
            row_errors.append("bill number is missing")
        if not admission_date:
            row_errors.append("admission date is missing or not recognized")
        if row_errors:
            errors.append((row_number, "; ".join(row_errors)))
            continue
        rows.append({
            "bill_number": bill_number,
            "patient_name": str(cell(row, "patient_name") or "").strip(),
            "patient_address": str(cell(row, "patient_address") or "").strip(),
            "admission_date": admission_date,
            "discharge_date": _parse_date_cell(cell(row, "discharge_date")),
            "consulting_doctor": str(cell(row, "consulting_doctor") or "").strip(),
            "referred_by": str(cell(row, "referred_by") or "").strip(),
            "final_amount": _parse_decimal_cell(cell(row, "final_amount")),
        })
    return rows, errors


def build_initial_rows(
    headers: list[str], data_rows: list[list], mapping: dict[str, str | None]
) -> tuple[list[dict], list[list[str]]]:
    """Apply `mapping` (target field -> source header) to every parsed row,
    producing one dict per row shaped like SalesEntryForm's initial data —
    ready to feed straight into a SalesEntryFormSet for review — plus a
    parallel list of human-readable warnings per row (unrecognized date,
    a category/customer name that doesn't match anything active, a zero
    amount) so the review page can flag exactly which rows need a manual
    fix before saving. Unmatched names are left blank rather than guessed
    at or auto-created."""
    index = {h: i for i, h in enumerate(headers)}

    def cell(row, field):
        header = mapping.get(field)
        if not header or header not in index:
            return ""
        i = index[header]
        return row[i] if i < len(row) else ""

    channels = list(Category.objects.filter(kind=Category.Kind.SALES, is_active=True))
    subcategories = list(Subcategory.objects.filter(is_active=True, category__kind=Category.Kind.SALES))
    customers = list(Customer.objects.filter(is_active=True))
    banks = list(BankAccount.objects.filter(is_active=True))

    initial_rows = []
    warnings = []
    for row in data_rows:
        row_warnings = []

        date_value = _parse_date_cell(cell(row, "date"))
        if not date_value:
            row_warnings.append("Date not recognized")

        raw_channel = str(cell(row, "channel") or "").strip()
        channel_id = _match_by_name(raw_channel, channels)
        if raw_channel and not channel_id:
            row_warnings.append(f"Category “{raw_channel}” not found")

        raw_subcategory = str(cell(row, "subcategory") or "").strip()
        subcategory_id = _match_by_name(raw_subcategory, subcategories)
        if raw_subcategory and not subcategory_id:
            row_warnings.append(f"Sub-category “{raw_subcategory}” not found")

        raw_customer = str(cell(row, "customer") or "").strip()
        customer_id = _match_by_name(raw_customer, customers)
        if raw_customer and not customer_id:
            row_warnings.append(f"Customer “{raw_customer}” not found")

        gross_amount = _parse_decimal_cell(cell(row, "gross_amount"))
        if gross_amount <= 0:
            row_warnings.append("Gross amount is 0")

        raw_bank = str(cell(row, "bank_account") or "").strip()
        bank_id = _match_by_name(raw_bank, banks)
        if raw_bank and not bank_id:
            row_warnings.append(f"Bank “{raw_bank}” not found")

        initial_rows.append({
            "date": date_value,
            "channel": channel_id or None,
            "subcategory": subcategory_id or None,
            "customer": customer_id or None,
            "quantity": _parse_int_cell(cell(row, "quantity")),
            "gross_amount": gross_amount,
            "discount": _parse_decimal_cell(cell(row, "discount")),
            "payment_mode": _match_payment_mode(cell(row, "payment_mode")) or PaymentMode.CASH,
            "bank_account": bank_id or None,
            "note": str(cell(row, "note") or "").strip(),
        })
        warnings.append(row_warnings)
    return initial_rows, warnings
