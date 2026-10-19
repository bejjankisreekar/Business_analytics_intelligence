"""Shared plumbing for the console's list pages: sortable columns,
pagination and CSV/Excel export of exactly what the filters show.

A list view calls `sorted_qs()` then either returns `export_response()`
(when `?export=csv|xlsx`) or `paginate()` for the page. Templates use
the `sa_sort` tag (superadmin_tags) for column headers and
superadmin/_pager.html for the page controls.
"""
import csv
import datetime
from decimal import Decimal

from django.core.paginator import Paginator
from django.http import HttpResponse
from django.utils import timezone

PAGE_SIZES = (25, 50, 100)
EXPORT_FORMATS = ("csv", "xlsx")


def sorted_qs(request, qs, allowed: dict, default: str):
    """Order `qs` by ?sort=<key> or ?sort=-<key>, where `allowed` maps a
    column key to an ORM field (or a tuple of them). Returns (qs, sort)."""
    sort = request.GET.get("sort", default)
    key = sort.lstrip("-")
    if key not in allowed:
        sort, key = default, default.lstrip("-")
    fields = allowed[key]
    fields = fields if isinstance(fields, tuple) else (fields,)
    desc = sort.startswith("-")
    from django.db.models import F

    ordering = [F(f).desc(nulls_last=True) if desc else F(f).asc(nulls_last=True) for f in fields]
    return qs.order_by(*ordering), sort


def page_numbers(page) -> list:
    """1 … 4 5 [6] 7 8 … 20"""
    total, current = page.paginator.num_pages, page.number
    shown = {1, total} | set(range(max(1, current - 2), min(total, current + 2) + 1))
    out, prev = [], 0
    for n in sorted(shown):
        if n - prev > 1:
            out.append("…")
        out.append(n)
        prev = n
    return out


def paginate(request, items, context: dict, default_size=25):
    """Adds page, page_numbers, page_query, page_size(s) to `context`;
    returns the page."""
    try:
        size = int(request.GET.get("per_page", default_size))
    except ValueError:
        size = default_size
    if size not in PAGE_SIZES:
        size = default_size
    page = Paginator(items, size).get_page(request.GET.get("page"))
    params = request.GET.copy()
    params.pop("page", None)
    context.update({
        "page": page,
        "page_numbers": page_numbers(page),
        "page_query": params.urlencode(),
        "page_size": size,
        "page_sizes": PAGE_SIZES,
        "export_query": _without(request, "page", "export"),
    })
    return page


def _without(request, *keys) -> str:
    params = request.GET.copy()
    for k in keys:
        params.pop(k, None)
    return params.urlencode()


def wants_export(request) -> str | None:
    fmt = request.GET.get("export")
    return fmt if fmt in EXPORT_FORMATS else None


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def export_response(fmt: str, filename: str, headers: list[str], rows):
    """A CSV or Excel download. `rows` is any iterable of sequences."""
    stamp = timezone.localdate().isoformat()
    if fmt == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}-{stamp}.csv"'
        response.write("﻿")  # so Excel opens UTF-8 (₹, names) correctly
        writer = csv.writer(response)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(["" if v is None else v for v in (_cell(v) for v in row)])
        return response

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = filename[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    widths = [len(h) + 2 for h in headers]
    for row in rows:
        values = [_cell(v) for v in row]
        ws.append(values)
        for i, v in enumerate(values):
            widths[i] = max(widths[i], min(48, len(str(v)) + 2))
            if isinstance(v, float):
                ws.cell(row=ws.max_row, column=i + 1).number_format = "#,##0.00"
            elif isinstance(v, (datetime.date, datetime.datetime)):
                ws.cell(row=ws.max_row, column=i + 1).number_format = (
                    "dd mmm yyyy hh:mm" if isinstance(v, datetime.datetime) else "dd mmm yyyy"
                )
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, w)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}-{stamp}.xlsx"'
    wb.save(response)
    return response
